# VecForge — Universal Local-First Vector Database
# Copyright (c) 2026 Suneel Bose K · ArcGX TechLabs Private Limited
# Licensed under the Business Source License 1.1 (BSL 1.1)

"""
Build the VecForge 2.0 paper as PDF and Word (.docx) from vecforge-paper.md.

    python paper/build_paper.py

Needs pandoc (3.x), Node with the `playwright` package and a Chromium build,
pypdf and reportlab. Math is typeset with KaTeX: set KATEX_DIR to a local
katex/dist folder, or leave it unset to load KaTeX 0.16.11 from jsDelivr.
Outputs paper/VecForge-2.0-Paper.pdf and paper/VecForge-2.0-Paper.docx.
"""

from __future__ import annotations

import io
import os
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
SRC = HERE / "vecforge-paper.md"
PDF = HERE / "VecForge-2.0-Paper.pdf"
DOCX = HERE / "VecForge-2.0-Paper.docx"
COMMON = [
    "--citeproc",
    "--bibliography",
    str(HERE / "references.bib"),
    "--csl",
    str(HERE / "ieee-numeric.csl"),
    "--number-sections",
    "--resource-path",
    str(HERE),
]
KATEX = os.environ.get("KATEX_DIR", "https://cdn.jsdelivr.net/npm/katex@0.16.11/dist")
# why: the page-number stamp must use an embedded font, or the PDF carries a non-embedded Times
ORCID_URL = "https://orcid.org/0009-0000-9282-2287"
FOOTER_FONT = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
pdfmetrics.registerFont(TTFont("LiberationSerif", FOOTER_FONT))


def build_html(out: Path) -> None:
    text = SRC.read_text(encoding="utf-8")
    abstract = re.search(r"abstract: \|\n((?:  .*\n)+)", text)
    plain = " ".join(abstract.group(1).split()) if abstract else ""
    subprocess.run(
        [
            "pandoc",
            str(SRC),
            "-o",
            str(out),
            "--standalone",
            "--template",
            str(HERE / "template.html"),
            f"--katex={KATEX.rstrip('/')}/",
            "-M",
            f"abstract-text={plain[:300]}",
            *COMMON,
        ],
        check=True,
    )


def footer(width: float, height: float, n: int, total: int) -> PdfReader:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height), initialFontName="LiberationSerif")
    c.setFillColor(HexColor("#666666"))
    c.setFont("LiberationSerif", 8.5)
    c.drawCentredString(width / 2, 30, f"{n}")
    c.save()
    buf.seek(0)
    return PdfReader(buf)


def build_pdf() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        html = HERE / "_paper.html"
        build_html(html)
        raw = Path(tmp) / "raw.pdf"
        # why: a real page session waits for KaTeX and its fonts before printing
        subprocess.run(
            ["node", str(HERE / "print_pdf.cjs"), html.as_uri(), str(raw)], check=True
        )
        html.unlink()
        reader = PdfReader(str(raw))
        total = len(reader.pages)
        writer = PdfWriter()
        for i, page in enumerate(reader.pages):
            page.merge_page(
                footer(
                    float(page.mediabox.width),
                    float(page.mediabox.height),
                    i + 1,
                    total,
                ).pages[0]
            )
            writer.add_page(page)
        writer.add_metadata(
            {
                "/Title": "VecForge: A Local-First Hybrid Vector Database with Persistent Agent Memory and Order-Preserving Grover-Inspired Reranking",
                "/Author": "Suneel Bose K (ORCID 0009-0000-9282-2287)",
                "/Subject": "Preprint, VecForge 2.0.0",
                "/Keywords": "vector database; hybrid retrieval; BM25; FAISS; agent memory; Model Context Protocol; Grover search; amplitude amplification",
                "/Creator": "paper/build_paper.py",
            }
        )
        with open(PDF, "wb") as f:
            writer.write(f)


def build_docx() -> None:
    ref = HERE / "reference.docx"
    subprocess.run(
        [
            "pandoc",
            str(SRC),
            "-o",
            str(DOCX),
            "--reference-doc",
            str(ref),
            "--lua-filter",
            str(HERE / "docx-keywords.lua"),
            *COMMON,
            "-M",
            f"date=ArcGX TechLabs Private Limited, Tiruchirappalli, Tamil Nadu, India. suneelbose@arcgx.in. ORCID iD {ORCID_URL}. October 2026",
        ],
        check=True,
    )
    finish_docx(DOCX)


def finish_docx(path: Path) -> None:
    """Stop Word from splitting a table row across pages, and link the ORCID iD on the title page."""
    doc = Document(str(path))
    for table in doc.tables:
        for row in table.rows:
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.append(OxmlElement("w:cantSplit"))
    for para in doc.paragraphs:
        if ORCID_URL in para.text and len(para.runs) == 1:
            before, after = para.runs[0].text.split(ORCID_URL, 1)
            para.runs[0].text = before
            rid = para.part.relate_to(ORCID_URL, RT.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rid)
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            rpr.append(style)
            run.append(rpr)
            text = OxmlElement("w:t")
            text.text = ORCID_URL
            run.append(text)
            link.append(run)
            para._p.append(link)
            para.add_run(after)
            break
    doc.save(str(path))


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print(f"wrote {PDF.name} ({len(PdfReader(str(PDF)).pages)} pages) and {DOCX.name}")
